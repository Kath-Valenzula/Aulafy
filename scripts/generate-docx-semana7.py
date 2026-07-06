#!/usr/bin/env python3
"""Genera el DOCX oficial Semana 7 desde la plantilla Duoc UC."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(
    "/Users/sbriceno/Documents/DUOC/SOFTWARE/TSY2201_EXP2_S7_Formato_respuesta_Finalizando el desarrollo de la solución.docx"
)
OUT = ROOT / "docs/semana-7/TSY2201_EXP2_S7_Formato_respuesta_Aulafy_Semana7.docx"
EVID = ROOT / "docs/semana-7/evidencias"


def add_paragraphs(cell, *blocks: str) -> None:
    cell.text = ""
    for i, block in enumerate(blocks):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.text = block


def main() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Plantilla no encontrada: {TEMPLATE}")

    shutil.copy(TEMPLATE, OUT)
    doc = Document(str(OUT))

    info = doc.tables[0]
    info.rows[0].cells[1].text = "Katherine Gisselle Valenzuela Moreno\nSebastián Alberto Briceño Inostroza"
    info.rows[1].cells[0].text = "Asignatura:\nTaller Aplicado de Software"
    info.rows[1].cells[1].text = "Carrera:\nIngeniería en Desarrollo de Software"
    info.rows[2].cells[0].text = "Profesor:\nAlonso Esteban Castillo Pizarro"
    info.rows[2].cells[1].text = "Fecha:\n5 de julio de 2026"

    add_paragraphs(
        doc.tables[1].rows[1].cells[0],
        "Repositorio GitHub: https://github.com/Kath-Valenzula/Aulafy",
        "Rama de entrega: develop (integración final en rama semana-7-new).",
        "Documentación Semana 7: docs/semana-7/ (plan, resultados, evidencias y DOCX oficial).",
    )

    add_paragraphs(
        doc.tables[2].rows[1].cells[0],
        "Frontend (Vercel — versión actual): https://aulafy-web.vercel.app",
        "Backend AWS Elastic Beanstalk: http://aulafy-api-staging.eba-uuqbidym.us-east-2.elasticbeanstalk.com/api",
        "Health check: http://aulafy-api-staging.eba-uuqbidym.us-east-2.elasticbeanstalk.com/api/health",
        "Vercel consume el backend AWS mediante proxy /api. Base de datos: Amazon RDS MySQL 8.x.",
    )

    avance_blocks = [
        "1. ESTADO GENERAL DEL AVANCE",
        "Aulafy cierra el MVP en Semana 7 con pruebas finales (caja blanca, negra y gris), documentación completa y despliegue accesible en nube.",
        "",
        "2. FUNCIONALIDADES IMPLEMENTADAS",
        "Autenticación JWT; dashboards por rol; muro; calendario; notas; asistencia; anotaciones; chat REST; riesgo académico; notificaciones Telegram opcionales.",
        "Chat en MVP para PROFESOR, APODERADO y ESTUDIANTE. Riesgo académico para ADMIN, COLEGIO y PROFESOR (alertas de sus cursos).",
        "Dashboard apoderado/estudiante incluye acceso directo a Mensajes del curso.",
        "",
        "3. ALCANCE ACORDADO CON EL DOCENTE",
        "Chat: DENTRO del MVP (Opción B). Sin WebSocket; recarga manual.",
        "Profesor jefe: role_in_course en BD (HEAD_TEACHER, SUBJECT_TEACHER, ASSISTANT). Permisos parciales en anotaciones y calendario según rol docente.",
        "Riesgo académico: promedio < 4.0 o asistencia < 85%. ADMIN/COLEGIO ven reporte institucional; PROFESOR ve alertas de sus cursos.",
        "",
        "4. VALIDACIONES Y PRUEBAS",
        "Caja blanca: 10 suites, 56 tests OK. Cobertura: Statements 64.87%, Branches 43.52%, Functions 52.21%, Lines 63.90%.",
        "Caja negra/gris: 17 casos manuales OK documentados en matriz. Evidencias local, AWS y Vercel.",
        "Build backend y frontend: OK. Fix UI Angular 21 con ChangeDetectorRef en chat, riesgo, usuarios y login.",
        "",
        "5. DOCUMENTOS Y EVIDENCIAS",
        "docs/semana-7/ con MD, DOCX, RESULTADOS_PRUEBAS_LOCAL.md y carpeta evidencias/ (local, aws, vercel).",
        "",
        "6. CUENTAS DEMO",
        "admin@aulafy.cl / Admin1234 | colegio@aulafy.cl / Colegio1234 | profesor@aulafy.cl / Profesor1234 | apoderado@aulafy.cl / Apoderado1234 | estudiante@aulafy.cl / Estudiante1234",
        "",
        "7. LIMITACIONES",
        "Chat sin tiempo real; moderación de chat no implementada; redeploy frontend S3 pendiente de credenciales AWS.",
    ]

    add_paragraphs(doc.tables[3].rows[1].cells[0], *avance_blocks)

    doc.add_page_break()
    doc.add_paragraph("EVIDENCIAS — CAPTURAS SEMANA 7")
    images = [
        ("1_GitHub_rama_semana7_new.png", "GitHub — rama semana-7-new"),
        ("6_Cobertura_tests_backend.png", "Cobertura Jest — 56 tests"),
        ("7_Matriz_pruebas_negras.png", "Matriz pruebas caja negra y gris"),
        ("8_Build_frontend_OK.png", "Build frontend exitoso"),
        ("vercel/9_Vercel_profesor_riesgo_academico.png", "Riesgo académico — rol PROFESOR"),
        ("vercel/10_Vercel_apoderado_dashboard_mensajes.png", "Dashboard apoderado — Mensajes del curso"),
        ("vercel/4_Vercel_admin_chat_funcionando.png", "Chat — Vercel + Beanstalk"),
        ("vercel/1_Vercel_proxy_health_Beanstalk.png", "Health backend vía proxy Vercel"),
    ]
    for filename, caption in images:
        img_path = EVID / filename
        if img_path.exists():
            doc.add_paragraph(caption)
            doc.add_picture(str(img_path), width=Inches(6.2))
            doc.add_paragraph("")

    doc.save(str(OUT))
    print(f"DOCX generado: {OUT}")


if __name__ == "__main__":
    main()
